from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.2)

st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    h = d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size)
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F); h.font.bold = True

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    tcPr.append(s)

def H1(t): d.add_heading(t, 1)
def H2(t): d.add_heading(t, 2)
def P(t, bold=False, italic=False, align=None):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if align: p.alignment = align
    return p
def B(items):
    for t in items: d.add_paragraph(t, style="List Bullet")
def N(items):
    for t in items: d.add_paragraph(t, style="List Number")
def N(items):
    for t in items: d.add_paragraph(t, style="List Number")
def CODE(t):
    p = d.add_paragraph(); r = p.add_run(t); r.font.name = "Consolas"; r.font.size = Pt(9)
    p.paragraph_format.left_indent = Cm(0.5)
def T(header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.size = Pt(10); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t

# ---------------- title page ----------------
for _ in range(3): d.add_paragraph()
P("PROJECT-BASED LEARNING", True, align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(14)
p = P("Intelligent Digital Library and Book Recommendation System", True, align=WD_ALIGN_PARAGRAPH.CENTER)
p.runs[0].font.size = Pt(24); p.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
P("Project Report (Phase-I design + working prototype)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
d.add_paragraph()
P("Team ID: DSCPP-II-2026-T367", True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Department of Computer Science & Engineering", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Graphic Era (Deemed to be University), Dehradun", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Academic Session 2026-27  |  Semester 3", align=WD_ALIGN_PARAGRAPH.CENTER)
d.add_paragraph()
P("Mentor: Dr. Jyoti Agarwal", True, align=WD_ALIGN_PARAGRAPH.CENTER)
d.add_paragraph()
T(["Team Member", "Roll No.", "Role"],
  [["Member 1: ____________________", "__________", "Core Engine (C) + Documentation"],
   ["Member 2: ____________________", "__________", "Search & Algorithms (C) + Main program + Testing"],
   ["Member 3: ____________________", "__________", "OOP & Recommendation (C++)"]], [6.5, 3.5, 6.5])
d.add_page_break()

# ---------------- 1 ----------------
H1("1. Introduction")
H2("1.1 Problem statement")
P("Most college and community libraries still use paper registers or a simple search box that only matches exact titles. "
  "A reader has to already know what they want. In a library with thousands of titles there is no easy way to search "
  "or to get suggestions, so readers browse blindly.")
H2("1.2 Motivation")
B(["Thousands of titles, but no easy way to search or get suggestions.",
   "Existing systems match exact titles only.",
   "Small libraries cannot afford expensive commercial software."])
H2("1.3 Who benefits")
T(["Group", "Benefit"],
  [["Students / readers (primary users)", "Find books faster, borrow and return easily, get suggestions they may like."],
   ["Librarians / admins (secondary)", "Add and remove books, see who has which book, create student accounts."]], [5, 11.5])

H1("2. Objectives")
B(["Build a digital library system that stores, searches, issues and returns books.",
   "Implement core operations (add, delete, search, issue, return) in C using data structures.",
   "Design an object-oriented model in C++ for users, books and recommendation logic.",
   "Build a recommendation feature based on genre, author and past reading activity.",
   "Test the complete system and demonstrate a working prototype with full documentation."])

H1("3. Technology used")
T(["Item", "Choice", "Why (plain words)"],
  [["Languages", "C and C++", "C for the fast book operations; C++ for classes and recommendation."],
   ["Storage", "Flat text files (data/*.txt)", "No database needed; easy to open and check in Notepad."],
   ["Compiler", "GCC (gcc + g++)", "Free and standard; compiles C and C++ together."],
   ["IDE", "Code::Blocks", "Easy for students; any project file can be added."],
   ["Interface", "Console menu", "Simple text menu, no graphics library required."]], [3, 4.5, 9])

H1("4. System architecture")
P("The program has four layers. The user talks only to the menu; everything else happens behind it.")
T(["User Interface\n(console menu)", "Library Core Engine\n(C, data structures)", "Recommendation Engine\n(C++, OOP)", "Data Storage\n(text files)"],
  [["Login, Student menu, Admin menu", "Linked list, hashing, BST, sorting, issue/return", "Genre / Author / Popularity rules", "books.txt, users.txt, issues.txt"]],
  [4, 4.2, 4.2, 4.1])
H2("4.1 Data / control flow")
N(["Program starts: books.txt is read into a linked list; a hash table and a BST are built from it.",
   "User logs in (users.txt). A Student or Admin object is created.",
   "Menu actions call the C core (search, issue, return, add, delete).",
   "'Recommend' asks the C++ Recommender, which reads the user's history and scores every unread book.",
   "Every change is written back to the text files immediately."])
H2("4.2 How C and C++ work together")
P("The C files (list.c, search.c) are compiled with gcc and the C++ files (oop.cpp, main.cpp) with g++. "
  "They share one header, library.h. The header is wrapped in extern \"C\" { } so that C++ code can call the C functions. "
  "This is the only 'trick' in the project.")

H1("5. Data structures and concepts used")
P("Each concept is explained in simple words, with where it works and where it does not.")
T(["Concept", "Where used", "Simple meaning", "Works well when", "Does NOT work well when"],
  [["Linked list", "list.c - stores all books", "Books chained one after another; each points to the next.",
    "Books are added/deleted often (no shifting). Example: admin deletes book 601 - just re-link nodes.",
    "You need to jump to the n-th item directly. Adding at the end walks the whole list (slow for 100,000 books)."],
   ["Hashing (chaining)", "search.c - find by book ID", "Book ID is turned into a table slot number (id % 101), so we jump straight to it.",
    "Looking up by exact unique ID. Example: search ID 105 is almost instant.",
    "You want a range (IDs 100-200) or partial matches; hashing cannot do that."],
   ["Binary Search Tree", "search.c - search by title", "Titles are kept in A-Z order in a tree; smaller goes left, bigger goes right.",
    "Finding a title quickly and listing titles A-Z. Example: 'harry' lists both Harry Potter books alphabetically.",
    "Titles are inserted already sorted - the tree becomes a long chain and slows to linear. Partial search still visits every node."],
   ["Sorting (insertion sort)", "search.c - sort by title / author / popularity", "Takes each book and inserts it at the right place in an already-sorted part.",
    "Small and medium lists (hundreds of books); simple to understand.",
    "Very large lists - it takes about n x n steps; a faster sort (merge/quick) would be needed."],
   ["File handling", "list.c, oop.cpp", "Data is saved in text files so it stays after the program closes.",
    "Small data, easy to inspect and back up.",
    "Many users at the same time, or very big data - no locking, whole file is rewritten."],
   ["Classes", "oop.h / oop.cpp", "A class bundles data and the functions that work on it (User, Student, Admin, Rule...).",
    "Keeping user logic and recommendation logic organised.", "Tiny programs where a plain function is enough."],
   ["Inheritance", "Student, Admin from User; GenreRule, AuthorRule, PopularityRule from Rule", "A child class reuses what its parent already has.",
    "Common features (search, sort) written once in User and used by both roles.", "When child and parent are not really 'the same kind of thing'."],
   ["Polymorphism", "User::menu(), Rule::score()", "Same function call, different behaviour depending on the real object.",
    "u->menu() shows the Student menu or Admin menu automatically; Recommender adds up any number of rules without knowing which.",
    "Adds a small overhead and is confusing if there are too many levels."]],
  [2.2, 2.7, 3.6, 4.0, 4.0])

H1("6. Module design (source files)")
T(["File", "Language", "What it contains"],
  [["src/library.h", "C / C++ shared", "Book structure and all C function declarations (with extern \"C\")."],
   ["src/list.c", "C", "Load/save books file, add, delete, display, issue, return, history, issued-books report."],
   ["src/search.c", "C", "Hash table (by ID), BST (by title), insertion sort (title/author/popularity)."],
   ["src/oop.h, src/oop.cpp", "C++", "User, Student, Admin classes; Rule classes; Recommender; login; input helpers."],
   ["src/main.cpp", "C++", "Login screen and main loop."],
   ["Makefile", "-", "One command (make) builds everything."]], [4, 3, 9.5])

H1("7. File formats")
P("All fields are separated by the | character (so the | character is not allowed inside titles).")
T(["File", "One line looks like", "Meaning"],
  [["data/books.txt", "201|Harry Potter ...|J K Rowling|Fantasy|1|3", "id | title | author | genre | available (1/0) | times issued"],
   ["data/users.txt", "priya|priya123|student", "username | password | role (admin/student)"],
   ["data/issues.txt", "priya|201|RETURNED", "username | book id | ISSUED or RETURNED (this is the reading history)"]], [3.3, 6.2, 7])

H1("8. OOP design")
CODE("User (abstract: getRole(), menu() are virtual)\n"
     "  |-- Student   : search, sort, issue, return, history, recommend\n"
     "  |-- Admin     : add/delete book, issued report, create student account\n\n"
     "Rule (abstract: score(book) is virtual)\n"
     "  |-- GenreRule      : +2 for each past book of the same genre\n"
     "  |-- AuthorRule     : +3 for each past book of the same author\n"
     "  |-- PopularityRule : + number of times the book was issued\n\n"
     "Recommender : creates the three rules, scores every unread book, shows top 5")
P("Role-based access: after login, the program creates a Student or an Admin object. The menu() call is the same line of code "
  "for both, but each class shows its own options - this is polymorphism. A student never sees admin options.")

H1("9. Recommendation algorithm")
H2("9.1 Steps")
B(["Read all book IDs the user has borrowed before (issues.txt) - this is the 'history'.",
   "For every book the user has NOT read: score = GenreRule + AuthorRule + PopularityRule.",
   "Sort by score (highest first) and show the top 5.",
   "If the user has no history, only popularity counts, so the most borrowed books are shown."])
H2("9.2 Worked example (user 'priya')")
P("History: 201 Harry Potter 1 (Fantasy, J K Rowling), 203 The Hobbit (Fantasy, J R R Tolkien), 102 Data Structures Using C (Programming).")
T(["Book", "Genre points", "Author points", "Popularity", "Total"],
  [["202 Harry Potter 2", "2 + 2 = 4 (two Fantasy books read)", "3 (Rowling read)", "1", "8"],
   ["204 The Lord of the Rings", "4", "3 (Tolkien read)", "0", "7"],
   ["101 The C Programming Language", "2 (one Programming book)", "0", "3", "5"],
   ["103 Let Us C", "2", "0", "2", "4"],
   ["205 The Alchemist", "0", "0", "3", "3"]], [5, 4.5, 3, 2, 2])
P("The program output matches this table exactly (Test T6 below).")
H2("9.3 Where it works and where it does not")
B(["Works: a reader who has borrowed several books of the same genre/author. Example: fantasy fan gets more fantasy.",
   "Works: a brand-new reader - gets popular books (a 'cold start' fallback).",
   "Does not work well: a reader with only one or two borrowed books - suggestions are weak.",
   "Does not work well: a library where all books have the same genre/author or no popularity data - scores tie.",
   "It does not understand book content (story, theme) - only genre, author and borrowing counts."])

H1("10. How to build and run")
B(["Linux / macOS / MinGW: open a terminal in the project folder and run  make  then  ./library  (Windows: library.exe).",
   "Code::Blocks: File > New > Empty project; add all files from src/; Build and Run. Set the working folder (Project > Properties > Build targets > Execution working dir) to the project folder so that data/ is found.",
   "Always start the program from the project folder, because the data files are opened as data/books.txt.",
   "Demo logins: admin / admin123 (Admin),  priya / priya123 and rahul / rahul123 (Students)."])

H1("11. Testing and results")
P("Tests were run by feeding scripted input to the real program (run_tests.sh). Each test starts from fresh sample data (25 books, 3 users). Full output is in docs/test_output.txt.")
T(["ID", "What was tested", "Expected", "Actual result", "Status"],
  [["T1", "Login with wrong password", "Rejected", "'Wrong username or password.'", "Pass"],
   ["T2", "Title search: 'the hobbit' (exact), 'harry' (partial), 'xyz'", "BST finds exact; partial gives A-Z list; not found message", "1 exact hit; 2 Harry Potter books in A-Z order; 'No book found'", "Pass"],
   ["T3", "Search by ID 105 and 999 (hashing)", "Found / not found", "Book 105 shown (Issued); 'No book with ID 999'", "Pass"],
   ["T4", "Sort by title, author, popularity", "Correct order", "Title starts 'A Brief History of Time'; author starts 'A P J Abdul Kalam'; popularity starts with books issued 3 times", "Pass"],
   ["T5", "Issue, issue again, wrong ID, return, return again", "Refuse invalid ones", "Issued OK; 'already issued'; 'No book'; Returned OK; 'not issued by you'", "Pass"],
   ["T6", "Recommendation for user with history", "Scores as in section 9.2", "202=8, 204=7, 101=5, 103=4, 205=3; already-read books not shown", "Pass"],
   ["T7", "Recommendation for brand-new student", "Most popular books", "Top 3 (score 3) then 2; message 'No reading history yet'", "Pass"],
   ["T8", "Admin: add book, duplicate ID, delete issued book, delete free book", "Add OK; refuse duplicate; refuse issued delete; delete OK", "All four behaved as expected", "Pass"],
   ["T9", "Admin: issued report; duplicate student username", "Report lists 2 rows; duplicate refused", "priya/102, rahul/105 listed; 'username already exists'", "Pass"],
   ["T10", "Data saved to files", "books.txt and issues.txt updated", "201 -> available 0, issued count 4; new ISSUED line added", "Pass"],
   ["Extra", "Invalid menu input (letters, wrong numbers)", "No crash", "'Invalid choice.' shown, program continues", "Pass"],
   ["Extra", "Memory check with valgrind (login, recommend, issue, return)", "No leaks / errors", "0 errors", "Pass"]],
  [1.2, 4.3, 3.6, 5.6, 1.5])

H1("12. Limitations and future scope")
T(["Limitation", "Why it matters", "Possible improvement"],
  [["Passwords stored as plain text", "Anyone who opens users.txt can read them.", "Store a hash of the password."],
   ["BST is not balanced", "Sorted input makes search slow.", "Use an AVL / balanced tree."],
   ["Insertion sort is slow for very big lists", "n x n steps.", "Use quick sort or merge sort."],
   ["Files are rewritten fully on every change", "Slow for huge data; not safe for many users at once.", "Use a database."],
   ["One copy per book", "Library cannot hold 2 copies of the same title.", "Add a 'copies' field."],
   ["No due dates or fines", "Real libraries need them.", "Store issue date and compute fine."],
   ["Recommendation uses only genre, author, popularity", "Cannot understand story or theme.", "Add keywords/ratings."],
   ["Console only", "Not user-friendly for everyone.", "Add a GUI or web front end."]], [4.5, 6, 6])

H1("13. Team work division (3 members)")
T(["Member", "Files / modules", "Concepts", "Report sections", "Testing duty"],
  [["Member 1\n(Core Engine + Docs)", "src/library.h, src/list.c, Makefile, data files", "Linked list, file handling, issue / return", "1, 2, 3, 4, 6, 7, 10, 12", "T5, T10"],
   ["Member 2\n(Search + Main + Testing)", "src/search.c, src/main.cpp, run_tests.sh", "Hashing, BST, sorting, program flow", "5 (hashing, BST, sorting), 11", "T1, T2, T3, T4, T8, T9"],
   ["Member 3\n(OOP + Recommendation)", "src/oop.h, src/oop.cpp", "Classes, inheritance, polymorphism, recommender, login", "5 (OOP rows), 8, 9", "T6, T7, T9"]], [3.2, 4, 3.6, 3.2, 2.5])
P("Each member should be able to explain their own files line by line (see docs/Viva_QnA.md) and give a short demo of their own part.")

H1("14. References")
B(["GeeksforGeeks - Data Structures & Algorithms tutorials",
   "GeeksforGeeks - Object Oriented Programming in C++",
   "cplusplus.com - C++ Standard Library reference",
   "Course lecture notes and lab material"])

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Intelligent Digital Library and Book Recommendation System  |  Page ")
r.font.size = Pt(9)
r2 = fp.add_run(); r2.font.size = Pt(9)
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    r2._r.append(e)

d.save("docs/Project_Report.docx")
print("saved")
